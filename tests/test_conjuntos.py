"""Testes de ponta a ponta com os comprovantes de exemplo e o gabarito."""
from pathlib import Path

import openpyxl
import pytest

from projeto_ir.avaliacao import CAMPOS, avaliar, carregar_gabarito
from projeto_ir.exportar import para_excel
from projeto_ir.processar import resumo
from projeto_ir.regras import carregar

TESTES = Path(__file__).resolve().parent
CONJUNTOS = [TESTES / "exemplos", TESTES / "exemplos_variados"]


@pytest.mark.parametrize("pasta", CONJUNTOS, ids=lambda p: p.name)
def test_acerto_no_gabarito(pasta):
    r, _ = avaliar(pasta)
    assert r.lidos > 0
    for campo in CAMPOS:
        assert r.taxa(campo) == 1.0, f"{campo}: " + "; ".join(e for e in r.erros if campo in e)
    assert r.precisao_alertas == 1.0 and r.cobertura_alertas == 1.0, "\n".join(r.erros)


def test_resumo_e_limite_de_educacao():
    pasta = TESTES / "exemplos_variados"
    contexto, _ = carregar_gabarito(pasta)
    regras = carregar(2026)
    _, comps = avaliar(pasta)
    r = resumo(comps, contexto, regras)
    dependente = contexto.dependentes[0].nome
    assert r["por_pessoa"][dependente]["educacao"] == regras.limite_educacao
    assert r["educacao_excedente"][dependente] == pytest.approx(11760.0 - regras.limite_educacao)


def test_reembolso_descontado():
    _, comps = avaliar(TESTES / "exemplos")
    consulta = next(c for c in comps if c.arquivo.startswith("04_"))
    assert consulta.valor_reembolsado == 300.0
    assert consulta.valor_dedutivel == 150.0


def test_planilha(tmp_path):
    pasta = TESTES / "exemplos"
    contexto, _ = carregar_gabarito(pasta)
    _, comps = avaliar(pasta)
    destino = para_excel(comps, contexto, carregar(2026), tmp_path / "saida.xlsx")
    wb = openpyxl.load_workbook(destino)
    assert wb.sheetnames == ["Resumo", "Comprovantes", "Alertas"]
    assert wb["Comprovantes"].max_row == len(comps) + 1
    assert wb["Alertas"].max_row > 1


def test_tela_abre_sem_erro():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(TESTES.parent / "projeto_ir" / "app.py"), default_timeout=60).run()
    assert not at.exception
    assert [m.label for m in at.metric] == ["Comprovantes", "Com alerta", "Total dedutível"]
