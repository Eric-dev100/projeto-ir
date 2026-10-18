"""Exporta a planilha de conferência para Excel."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .modelo import Comprovante, Contexto
from .processar import resumo
from .regras import Regras

NOMES_CATEGORIA = {"saude": "Saúde", "educacao": "Educação", "plano_saude": "Plano de saúde",
                   "nao_dedutivel": "Não dedutível", "indefinida": "Indefinida"}
MOEDA = 'R$ #,##0.00'
CABECALHO = PatternFill("solid", fgColor="DCEFE8")
ALERTA = PatternFill("solid", fgColor="FCEBD5")


def _cabecalho(ws, colunas: list[tuple[str, int]]) -> None:
    for i, (titulo, largura) in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=i, value=titulo)
        cel.font = Font(bold=True)
        cel.fill = CABECALHO
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"


def _data_br(iso: str | None) -> str:
    return "/".join(reversed(iso.split("-"))) if iso else ""


def para_excel(comps: list[Comprovante], contexto: Contexto, regras: Regras, destino: Path) -> Path:
    wb = Workbook()
    r = resumo(comps, contexto, regras)

    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Projeto IR - conferência do ano-calendário {contexto.ano_calendario}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Titular: {contexto.titular.nome}. A ferramenta organiza e alerta, mas não decide: "
                "confira tudo antes de declarar.")
    if regras.situacao != "revisado":
        ws["A3"] = "Atenção: as regras fiscais deste ano ainda não foram revisadas para a declaração atual."
        ws["A3"].font = Font(color="9C5700", bold=True)
    linha = 5
    for i, t in enumerate(["Pessoa", "Categoria", "Valor dedutível"], start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font, c.fill = Font(bold=True), CABECALHO
    for pessoa, cats in r["por_pessoa"].items():
        for cat, valor in sorted(cats.items()):
            linha += 1
            ws.cell(row=linha, column=1, value=pessoa)
            ws.cell(row=linha, column=2, value=NOMES_CATEGORIA.get(cat, cat))
            ws.cell(row=linha, column=3, value=round(valor, 2)).number_format = MOEDA
    linha += 2
    for pessoa, excedente in r["educacao_excedente"].items():
        ws.cell(row=linha, column=1, value=(f"{pessoa}: educação passou do limite anual em R$ {excedente:,.2f}; "
                                            "o excedente não pode ser deduzido."))
        linha += 1
    ws.cell(row=linha, column=1, value=f"Comprovantes analisados: {r['total_arquivos']} | com alerta: {r['com_alerta']}")
    for col, larg in zip("ABC", (34, 18, 18)):
        ws.column_dimensions[col].width = larg

    ws = wb.create_sheet("Comprovantes")
    _cabecalho(ws, [("Arquivo", 34), ("Emitente", 38), ("CPF/CNPJ", 20), ("Documento válido", 11),
                    ("Data", 12), ("Valor", 14), ("Reembolsado", 14), ("Valor dedutível", 15),
                    ("Beneficiário", 28), ("Categoria", 16), ("Alertas", 40)])
    for i, c in enumerate(comps, start=2):
        valores = [c.arquivo, c.emitente, c.documento_emitente,
                   {True: "Sim", False: "Não", None: ""}[c.documento_valido], _data_br(c.data),
                   c.valor, c.valor_reembolsado or None, c.valor_dedutivel,
                   c.beneficiario, NOMES_CATEGORIA.get(c.categoria, c.categoria),
                   ", ".join(sorted(c.codigos_alerta))]
        for j, v in enumerate(valores, start=1):
            cel = ws.cell(row=i, column=j, value=v)
            if j in (6, 7, 8):
                cel.number_format = MOEDA
            if c.alertas:
                cel.fill = ALERTA

    ws = wb.create_sheet("Alertas")
    _cabecalho(ws, [("Arquivo", 34), ("Alerta", 24), ("O que aconteceu", 60), ("O que fazer", 70)])
    linha = 2
    for c in comps:
        for a in c.alertas:
            for j, v in enumerate([c.arquivo, a.codigo, a.mensagem, a.sugestao], start=1):
                ws.cell(row=linha, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            linha += 1

    destino = Path(destino)
    wb.save(destino)
    return destino
